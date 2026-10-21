"""Бюджет дома «под чистовую отделку»: договор ООО «Строй Хауз» (фундамент, стены, кровля) +
крупные блоки работ до чистовой отделки. Свет — Maytoni. Суммы в ₽, вывод в млн ₽."""
import csv
import io
import json
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
import house_data as H
import make_rd_budget as R

OUT = sys.argv[1] if len(sys.argv) > 1 else "."
SCRATCH = os.environ.get("SCRATCH", "/tmp")

CONTRACT = 42_000_000      # договор Строй Хауз: фундамент, стены, кровля
RESERVE = 0.07             # непредвиденные — на всё, кроме фиксированного договора
LIGHT = ("Светильник",)    # позиции ЭОМ, которые заменяет блок Maytoni

FLOOR = round(sum(r["area"] for r in H.heat_losses()))
A_WIN = round(sum(r["A_win"] for r in H.heat_losses()))

# Свет Maytoni: (статья, кол-во, ед., цена мин., цена макс.), ₽
MAYTONI = [
    ("Встраиваемые и накладные споты Maytoni Technical", 75, "шт.", 3500, 9000),
    ("Магнитный трек Maytoni (шинопровод с аксессуарами)", 40, "м", 3000, 6000),
    ("Трековые светильники Maytoni", 25, "шт.", 7000, 16000),
    ("Подвесные и настенные декоративные светильники Maytoni", 12, "шт.", 15000, 60000),
    ("Люстры второго света: холл и гостиная", 3, "шт.", 80000, 250000),
    ("Светодиодный профиль: ниши, лестница, карнизы", 60, "м", 3000, 6000),
    ("Фасад и терраса, Maytoni Outdoor IP65", 8, "шт.", 6000, 15000),
    ("Сауна и С/У: влагозащищённые IP44/IP54", 8, "шт.", 4000, 10000),
    ("Блоки питания, диммеры DALI / 0–10 В", 1, "компл.", 80000, 250000),
    ("Монтаж светильников, треков и профиля", 190, "точ.", 800, 1500),
    ("Светотехнический проект (расчёт освещённости, раскладка)", 1, "компл.", 80000, 200000),
]


def maytoni():
    lo = sum(q * a for _, q, _u, a, _b in MAYTONI)
    hi = sum(q * b for _, q, _u, _a, b in MAYTONI)
    return lo, hi


def rd(code, exclude=()):
    return R.totals(code, exclude)["total"]


# Блоки: (группа, название, мин., макс., пометка)
def blocks():
    sh = "если не входит в договор Строй Хауза"
    return [
        ("Коробка", "Строй Хауз: фундамент, стены, кровля (договор)", CONTRACT, CONTRACT, "фикс."),
        ("Коробка", f"Окна и витражи ≈ {A_WIN} м², энергосберегающие, с монтажом",
         A_WIN * 25000, A_WIN * 40000, ""),
        ("Коробка", "Входные двери", 300000, 800000, ""),
        ("Коробка", "Фасад: шлифовка, покраска бруса, подшивка свесов, цоколь", 1200000, 2200000, sh),
        ("Коробка", "Каркасные перегородки (КД6)", 800000, 1400000, sh),
        ("Коробка", f"Стяжки пола по тёплому полу ≈ {FLOOR} м²", FLOOR * 1300, FLOOR * 2200, sh),
        ("Коробка", "Лестница (каркас и ступени)", 800000, 2000000, ""),
        ("Инженерия", "Вентиляция и кондиционирование (ПВУ с рекуператором, VRF)", *rd("291-ОВК"), ""),
        ("Инженерия", "Отопление: пеллетный котёл, котельная, тёплые полы, конвекторы", *rd("291-ОВ"), ""),
        ("Инженерия", "Водопровод и канализация (без сантехприборов)", *rd("291-ВК"), ""),
        ("Инженерия", "Электрика: щиты, кабели, розетки, выключатели", *rd("291-ЭОМ", LIGHT), ""),
        ("Инженерия", "Свет Maytoni во всём доме", *maytoni(), "с монтажом и светопроектом"),
        ("Инженерия", "Умный дом, сеть, Wi-Fi, видеонаблюдение, домофон", *rd("291-СС"), ""),
        ("Инженерия", "Молниезащита, заземление, пожарная сигнализация", *rd("291-МЗ.ПС"), ""),
        ("Инженерия", "Наружные сети и подключения (вода, канализация, 50 кВт, дренаж)",
         *rd("291-НС"), "до ТУ"),
        ("Организация", "Проекты, топосъёмка, технадзор", 690000, 1590000, ""),
        ("Организация", "Временные сооружения, вывоз мусора, охрана", 400000, 800000, ""),
    ]


LATER = [  # не входит в «под чистовую»
    ("Чистовая отделка: полы, плитка, покраска, потолки, межкомнатные двери", 8800000, 19500000),
    ("Сантехприборы по дизайн-проекту", 1200000, 3500000),
    ("Газ, этап 2: котёл, газопровод по участку, проект", 540000, 1400000),
    ("Благоустройство: отмостка, дорожки, озеленение", 1300000, 4500000),
    ("КНС — только если нет самотёка (по ТУ)", 350000, 800000),
]


def summary():
    b = blocks()
    var_lo = sum(x[2] for x in b[1:])
    var_hi = sum(x[3] for x in b[1:])
    res = (var_lo * RESERVE, var_hi * RESERVE)
    tot = (CONTRACT + var_lo + res[0], CONTRACT + var_hi + res[1])
    sh_opt = [x for x in b if "Строй Хауза" in x[4]]
    return dict(blocks=b, reserve=res, total=tot,
                sh_opt=(sum(x[2] for x in sh_opt), sum(x[3] for x in sh_opt)),
                groups={g: (sum(x[2] for x in b if x[0] == g), sum(x[3] for x in b if x[0] == g))
                        for g in ("Коробка", "Инженерия", "Организация")},
                later=LATER, maytoni=MAYTONI, floor=FLOOR)


def m(x):
    return round(x / 1e6, 1)


def xlsx(path, s):
    wb = Workbook()
    ws = wb.active
    ws.title = "Под чистовую"
    thin = Side(style="thin")
    brd = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.append(["Группа", "Блок", "Мин., ₽", "Макс., ₽", "Среднее, ₽", "Примечание"])
    for g, n, lo, hi, note in s["blocks"]:
        r = ws.max_row + 1
        ws.append([g, n, round(lo, -3), round(hi, -3), f"=(C{r}+D{r})/2", note])
    last = ws.max_row
    r = ws.max_row + 1
    ws.append(["", f"Непредвиденные {RESERVE:.0%} (кроме договора Строй Хауз)",
               f"=ROUND({RESERVE}*SUM(C3:C{last}),-3)", f"=ROUND({RESERVE}*SUM(D3:D{last}),-3)",
               f"=(C{r}+D{r})/2", ""])
    t = ws.max_row + 1
    ws.append(["", "ИТОГО под чистовую", f"=SUM(C2:C{r})", f"=SUM(D2:D{r})", f"=(C{t}+D{t})/2", ""])
    ws.append([])
    ws.append(["", "После — не входит в итог", None, None, None, ""])
    for n, lo, hi in LATER:
        r2 = ws.max_row + 1
        ws.append(["Позже", n, lo, hi, f"=(C{r2}+D{r2})/2", ""])
    ws2 = wb.create_sheet("Свет Maytoni")
    ws2.append(["Статья", "Кол-во", "Ед.", "Цена мин., ₽", "Цена макс., ₽", "Сумма мин., ₽", "Сумма макс., ₽"])
    for n, q, u, a, b in MAYTONI:
        r3 = ws2.max_row + 1
        ws2.append([n, q, u, a, b, f"=B{r3}*D{r3}", f"=B{r3}*E{r3}"])
    r3 = ws2.max_row + 1
    ws2.append(["ИТОГО свет Maytoni", None, None, None, None, f"=SUM(F2:F{r3 - 1})", f"=SUM(G2:G{r3 - 1})"])
    for sh, widths in ((ws, [14, 70, 15, 15, 15, 34]), (ws2, [60, 9, 8, 14, 14, 15, 15])):
        for c in sh[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDDDDD")
        for i, wdt in enumerate(widths):
            sh.column_dimensions[chr(65 + i)].width = wdt
        for row in sh.iter_rows(min_row=1, max_row=sh.max_row):
            for c in row:
                c.border = brd
                c.alignment = Alignment(wrap_text=True, vertical="center")
                if isinstance(c.value, (int, float)) or (isinstance(c.value, str) and c.value.startswith("=")):
                    c.number_format = "# ##0"
    for c in ws[t]:
        c.font = Font(bold=True)
    wb.save(path)


def to_csv(s):
    b = io.StringIO()
    w = csv.writer(b)
    w.writerow(["Группа", "Блок", "Мин., млн ₽", "Макс., млн ₽", "Примечание"])
    for g, n, lo, hi, note in s["blocks"]:
        w.writerow([g, n, m(lo), m(hi), note])
    w.writerow(["", f"Непредвиденные {RESERVE:.0%}", m(s["reserve"][0]), m(s["reserve"][1]), ""])
    w.writerow(["", "ИТОГО под чистовую", m(s["total"][0]), m(s["total"][1]), ""])
    w.writerow([])
    for n, lo, hi in LATER:
        w.writerow(["Позже (не в итоге)", n, m(lo), m(hi), ""])
    return b.getvalue()


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    s = summary()
    xlsx(os.path.join(OUT, "Бюджет под чистовую — Берта Ромарио.xlsx"), s)
    open(os.path.join(SCRATCH, "prefinish.csv"), "w").write(to_csv(s))
    json.dump(s, open(os.path.join(SCRATCH, "prefinish.json"), "w"), ensure_ascii=False, indent=1)
    for g, n, lo, hi, note in s["blocks"]:
        print(f"{g:12} {n[:70]:70} {m(lo):6} {m(hi):6} {note}")
    print("резерв", m(s["reserve"][0]), m(s["reserve"][1]))
    print("ИТОГО", m(s["total"][0]), m(s["total"][1]), "группы", {k: (m(a), m(b)) for k, (a, b) in s["groups"].items()})
    print("опц. СХ", m(s["sh_opt"][0]), m(s["sh_opt"][1]))
